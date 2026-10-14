import hashlib

from pptx import Presentation

from conftest import EXAMPLES
from image2ppt import pipeline
from image2ppt.demo import draw_slide
from image2ppt.render import pptx_stats


def build(tmp_path, name="out.pptx"):
    tmp_path.mkdir(exist_ok=True)
    png = tmp_path / "slide.png"
    draw_slide(png)
    out = tmp_path / name
    layouts = pipeline.convert([png], out, layout=EXAMPLES / "demo_layout.raw.json", rounds=0,
                               workdir=tmp_path / "work", log=lambda m: None)
    return out, layouts[0]


def expected_counts(lay):
    els = lay["elements"]
    segments = sum(len(e["points"]) - 1 for e in els if e["type"] in ("arrow", "line"))
    return {
        "shapes": sum(e["type"] in ("rect", "text") for e in els),
        "text_boxes": sum(e["type"] == "text" for e in els),
        "pictures": sum(e["type"] == "crop" for e in els),
        "connectors": segments,
    }


def test_pptx_reopens_with_expected_shape_counts(tmp_path):
    out, lay = build(tmp_path)
    prs = Presentation(str(out))
    assert len(prs.slides) == 1
    st = pptx_stats(out)
    exp = expected_counts(lay)
    for key, value in exp.items():
        assert st[key] == value, key
    assert len(prs.slides[0].shapes) == st["shapes"] + st["pictures"] + st["connectors"]


def test_connector_ends_are_attached_to_boxes(tmp_path):
    out, lay = build(tmp_path)
    st = pptx_stats(out)
    two_point = [e for e in lay["elements"] if e["type"] == "arrow" and len(e["points"]) == 2]
    both = [e for e in two_point if e["anchors"].get("begin", {}).get("cxn") and e["anchors"].get("end", {}).get("cxn")]
    assert len(both) >= 3
    assert st["connectors_both_ends"] == len(both)
    # the connection ids must name existing shapes
    slide = Presentation(str(out)).slides[0]
    ids = {sh.shape_id for sh in slide.shapes}
    from pptx.oxml.ns import qn
    for cx in slide.shapes._spTree.iter(qn("p:cxnSp")):
        for tag in ("a:stCxn", "a:endCxn"):
            el = cx.find(qn("p:nvCxnSpPr")).find(qn("p:cNvCxnSpPr")).find(qn(tag))
            if el is not None:
                assert int(el.get("id")) in ids


def test_output_is_byte_reproducible(tmp_path):
    a, _ = build(tmp_path / "one")
    b, _ = build(tmp_path / "two")
    assert hashlib.sha256(a.read_bytes()).hexdigest() == hashlib.sha256(b.read_bytes()).hexdigest()
