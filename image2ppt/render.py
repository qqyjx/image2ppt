"""Layout dict -> editable PPTX.

Rects, text and lines become native shapes. Line ends that tidy anchored to a
box are attached to that box's connection point (so the line follows when the
box is dragged). 'crop' regions are cut from the source image and inserted as
pictures. The output file is byte-reproducible for the same input.
"""
import io
import zipfile

from PIL import Image
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

SLIDE_W_IN = 13.3333
CXN_IDX = {"top": 0, "left": 1, "bottom": 2, "right": 3}
FONT = "Arial"


def _rgb(s, default="1E1E1E"):
    s = (s or default).lstrip("#")
    return RGBColor.from_string(s.upper()[:6])


def _knockout(im, colors, tol=16):
    im = im.convert("RGBA")
    px = im.load()
    for yy in range(im.height):
        for xx in range(im.width):
            r, g, b, a = px[xx, yy]
            for kr, kg, kb in colors:
                if abs(r - kr) < tol and abs(g - kg) < tol and abs(b - kb) < tol:
                    px[xx, yy] = (r, g, b, 0)
                    break
    return im


def _arrowhead(connector):
    ln = connector.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")


def add_slide(prs, lay, src_img):
    W, H = lay["image_size"]
    sw = SLIDE_W_IN * 914400
    sh = sw * H / W
    prs.slide_width, prs.slide_height = Emu(int(sw)), Emu(int(sh))
    sx = lambda px: Emu(int(px / W * sw))
    sy = lambda px: Emu(int(px / H * sh))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = _rgb(lay.get("background"), "FFFFFF")
    bgc = tuple(int(lay.get("background", "FFFFFF").lstrip("#")[i:i + 2], 16) for i in (0, 2, 4))
    shapes = {}

    def set_text(shape, el):
        tf = shape.text_frame
        tf.word_wrap = el["type"] == "rect"  # free labels never re-wrap
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, Pt(2))
        for i, ln in enumerate(el["text"].split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT}.get(el.get("align"), PP_ALIGN.CENTER)
            r = p.add_run()
            r.text = ln
            r.font.size = Pt(el.get("size", 10))
            r.font.bold = bool(el.get("bold"))
            r.font.italic = bool(el.get("italic"))
            r.font.name = FONT
            r.font.color.rgb = _rgb(el.get("color"), "1E1E1E")

    # boxes first (the extractor lists parents before their children)
    for el in lay["elements"]:
        t = el["type"]
        if t == "rect":
            x, y, w, h = el["box"]
            kind = MSO_SHAPE.ROUNDED_RECTANGLE if el.get("rounded") else MSO_SHAPE.RECTANGLE
            shp = slide.shapes.add_shape(kind, sx(x), sy(y), sx(w), sy(h))
            if el.get("rounded"):
                radius_px = max(6.0, 0.008 * W)
                shp.adjustments[0] = min(0.5, radius_px / max(1.0, min(w, h)))
            if el.get("fill"):
                shp.fill.solid()
                shp.fill.fore_color.rgb = _rgb(el["fill"])
            else:
                shp.fill.background()
            if el.get("line"):
                shp.line.color.rgb = _rgb(el["line"])
                shp.line.width = Pt(1.25)
                if el.get("dash"):
                    shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            else:
                shp.line.fill.background()
            shp.shadow.inherit = False
            if el.get("text"):
                set_text(shp, el)
            shapes[el["id"]] = shp
        elif t == "text":
            x, y, w, h = el["box"]
            shp = slide.shapes.add_textbox(sx(x), sy(y), sx(w), sy(h))
            set_text(shp, el)
            shapes[el["id"]] = shp
        elif t == "crop":
            x, y, w, h = el["box"]
            m = 2  # small margin so tight crops do not clip edges
            x, y = max(0, x - m), max(0, y - m)
            w, h = min(W - x, w + 2 * m), min(H - y, h + 2 * m)
            im = src_img.crop((int(x), int(y), int(x + w), int(y + h)))
            if el.get("key_bg"):
                im = _knockout(im, [bgc])
            buf = io.BytesIO()
            im.save(buf, "PNG")
            buf.seek(0)
            pic = slide.shapes.add_picture(buf, sx(x), sy(y), sx(w), sy(h))
            shapes[el["id"]] = pic

    # lines last so they sit on top of the boxes they attach to
    for el in lay["elements"]:
        if el["type"] not in ("arrow", "line"):
            continue
        P = el.get("points") or [el["from"], el["to"]]
        anchors = el.get("anchors", {})
        for i in range(len(P) - 1):
            (x1, y1), (x2, y2) = P[i], P[i + 1]
            c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, sx(x1), sy(y1), sx(x2), sy(y2))
            c.line.color.rgb = _rgb(el.get("color"), "1E1E1E")
            c.line.width = Pt(el.get("width_pt", 1.5))
            if el.get("dash"):
                c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            # only the first and last segment of an elbowed path can attach to boxes
            if i == 0 and anchors.get("begin", {}).get("cxn") and anchors["begin"]["id"] in shapes:
                c.begin_connect(shapes[anchors["begin"]["id"]], CXN_IDX[anchors["begin"]["side"]])
            if i == len(P) - 2 and anchors.get("end", {}).get("cxn") and anchors["end"]["id"] in shapes:
                c.end_connect(shapes[anchors["end"]["id"]], CXN_IDX[anchors["end"]["side"]])
            if el["type"] == "arrow" and i == len(P) - 2:
                _arrowhead(c)
    return slide


def _save_reproducible(prs, out_path):
    """Save with fixed zip timestamps so the same layout always gives the same bytes."""
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            new = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            new.compress_type = zipfile.ZIP_DEFLATED
            new.external_attr = info.external_attr
            zout.writestr(new, zin.read(info.filename))


def render(pages, out_path):
    """pages: list of (layout_dict, PIL.Image). Writes one PPTX with one slide per page."""
    prs = Presentation()
    for lay, img in pages:
        add_slide(prs, lay, img.convert("RGB"))
    _save_reproducible(prs, str(out_path))
    return out_path


def pptx_stats(path):
    """Count what a PPTX contains, straight from the slide XML.

    Returns a dict with: slides, shapes (all p:sp, text boxes included), text_boxes,
    pictures, connectors, connectors_both_ends (a connection attached at the begin
    and at the end), connectors_one_end.
    """
    from pptx.oxml.ns import qn as _qn
    prs = Presentation(str(path))
    st = dict(slides=len(prs.slides), shapes=0, text_boxes=0, pictures=0,
              connectors=0, connectors_both_ends=0, connectors_one_end=0)
    for slide in prs.slides:
        tree = slide.shapes._spTree
        for sp in tree.iter(_qn("p:sp")):
            st["shapes"] += 1
            if sp.find(_qn("p:nvSpPr")).find(_qn("p:cNvSpPr")).get("txBox") == "1":
                st["text_boxes"] += 1
        st["pictures"] += len(list(tree.iter(_qn("p:pic"))))
        for cx in tree.iter(_qn("p:cxnSp")):
            st["connectors"] += 1
            c_nv = cx.find(_qn("p:nvCxnSpPr")).find(_qn("p:cNvCxnSpPr"))
            begin, end = c_nv.find(_qn("a:stCxn")) is not None, c_nv.find(_qn("a:endCxn")) is not None
            if begin and end:
                st["connectors_both_ends"] += 1
            elif begin or end:
                st["connectors_one_end"] += 1
    return st
